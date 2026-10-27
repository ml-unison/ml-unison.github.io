"""Valida y exporta el banco de preguntas del curso.

Fuente única: quizzes/banco/uNN.yml. Este script genera, en quizzes/exportados/:

* uNN.gift  — formato GIFT (Moodle y otras plataformas).
* uNN.xml   — Moodle XML (conserva retroalimentación por opción).
* uNN.docx  — documento para "Quick Import" de Microsoft Forms/Teams
              (las respuestas correctas van en una tabla al final, porque
              Forms no las importa y hay que marcarlas a mano).

Uso:
    uv run python herramientas/exportar_quizzes.py            # todos los bancos
    uv run python herramientas/exportar_quizzes.py u01 u02    # algunos
    uv run python herramientas/exportar_quizzes.py --validar  # solo validar
"""

from __future__ import annotations

import argparse
import html
import re
import sys
from pathlib import Path
from xml.etree import ElementTree as ET

import yaml

RAIZ = Path(__file__).resolve().parents[1]
BANCO = RAIZ / "quizzes" / "banco"
SALIDA = RAIZ / "quizzes" / "exportados"
TIPOS = {"opcion", "multiple", "vf", "numerica"}


class ErrorBanco(ValueError):
    pass


# --------------------------------------------------------------------------
# Carga y validación
# --------------------------------------------------------------------------
def cargar(nombre: str) -> dict:
    ruta = BANCO / f"{nombre}.yml"
    with ruta.open(encoding="utf-8") as f:
        banco = yaml.safe_load(f)
    validar(banco, nombre)
    return banco


def validar(banco: dict, nombre: str) -> None:
    errores: list[str] = []
    if not isinstance(banco, dict) or "preguntas" not in banco:
        raise ErrorBanco(f"{nombre}: falta la lista 'preguntas'")
    if not banco.get("titulo"):
        errores.append("falta 'titulo'")
    vistos: set[str] = set()
    for i, p in enumerate(banco["preguntas"], 1):
        pid = p.get("id") or f"#{i}"
        pref = f"{nombre}/{pid}"
        if not p.get("id"):
            errores.append(f"{pref}: falta 'id'")
        elif not str(p["id"]).startswith(nombre + "-"):
            errores.append(f"{pref}: el id debe empezar con '{nombre}-'")
        if pid in vistos:
            errores.append(f"{pref}: id repetido")
        vistos.add(pid)
        tipo = p.get("tipo", "opcion")
        if tipo not in TIPOS:
            errores.append(f"{pref}: tipo '{tipo}' no válido")
            continue
        if not str(p.get("enunciado", "")).strip():
            errores.append(f"{pref}: enunciado vacío")
        if tipo in {"opcion", "multiple"}:
            ops = p.get("opciones") or []
            correctas = sum(bool(o.get("correcta")) for o in ops)
            if len(ops) < 2:
                errores.append(f"{pref}: necesita al menos 2 opciones")
            if tipo == "opcion" and correctas != 1:
                errores.append(f"{pref}: 'opcion' requiere exactamente 1 correcta (tiene {correctas})")
            if tipo == "multiple" and correctas < 1:
                errores.append(f"{pref}: 'multiple' requiere al menos 1 correcta")
        elif tipo == "vf":
            if not isinstance(p.get("respuesta"), bool):
                errores.append(f"{pref}: 'vf' requiere respuesta true/false")
        elif tipo == "numerica":
            if not isinstance(p.get("respuesta"), (int, float)):
                errores.append(f"{pref}: 'numerica' requiere respuesta numérica")
            if not isinstance(p.get("tolerancia", 0), (int, float)):
                errores.append(f"{pref}: tolerancia debe ser numérica")
    if errores:
        raise ErrorBanco("\n".join(errores))


# --------------------------------------------------------------------------
# Utilidades de texto
# --------------------------------------------------------------------------
_MATH_BLOQUE = re.compile(r"\$\$(.+?)\$\$", re.S)
_MATH_LINEA = re.compile(r"(?<![\\$])\$(?!\$)(.+?)(?<!\\)\$")


def latex_moodle(texto: str) -> str:
    """$..$ → \\(..\\) y $$..$$ → \\[..\\] (filtro MathJax de Moodle)."""
    texto = _MATH_BLOQUE.sub(lambda m: r"\[" + m.group(1) + r"\]", texto)
    return _MATH_LINEA.sub(lambda m: r"\(" + m.group(1) + r"\)", texto)


def a_html(texto: str) -> str:
    """Conversión mínima Markdown → HTML (párrafos, `código`, **negritas**, *itálicas*)."""
    texto = latex_moodle(str(texto or "").strip())
    texto = html.escape(texto, quote=False)
    texto = re.sub(r"`([^`]+)`", r"<code>\1</code>", texto)
    texto = re.sub(r"\*\*([^*]+)\*\*", r"<strong>\1</strong>", texto)
    texto = re.sub(r"(?<!\*)\*([^*]+)\*(?!\*)", r"<em>\1</em>", texto)
    parrafos = [p.replace("\n", " ") for p in re.split(r"\n\s*\n", texto) if p.strip()]
    return "".join(f"<p>{p}</p>" for p in parrafos)


def plano(texto: str) -> str:
    """Texto plano de una línea (para DOCX)."""
    return re.sub(r"\s+", " ", str(texto or "")).strip()


_GIFT_ESPECIALES = re.compile(r"([~=#{}:\\])")


def gift_esc(texto: str) -> str:
    """Texto → HTML de una línea, con los caracteres especiales de GIFT escapados."""
    html_txt = a_html(str(texto or ""))
    if html_txt.count("<p>") == 1:  # un solo párrafo: sin envoltura
        html_txt = html_txt[3:-4]
    return _GIFT_ESPECIALES.sub(r"\\\1", plano(html_txt))


# --------------------------------------------------------------------------
# GIFT
# --------------------------------------------------------------------------
def a_gift(banco: dict) -> str:
    lineas = [f"// {banco['titulo']}", "// Generado por herramientas/exportar_quizzes.py — no editar a mano", ""]
    lineas.append(f"$CATEGORY: $course$/Reconocimiento de Patrones/{banco['titulo']}")
    lineas.append("")
    for p in banco["preguntas"]:
        tipo = p.get("tipo", "opcion")
        cab = f"::{p['id']}::[html]{gift_esc(p['enunciado'])}"
        expl = f"####{gift_esc(p['explicacion'])}" if p.get("explicacion") else ""
        if tipo == "vf":
            lineas.append(f"{cab} {{{'T' if p['respuesta'] else 'F'}{expl}}}")
        elif tipo == "numerica":
            tol = p.get("tolerancia", 0)
            lineas.append(f"{cab} {{#{p['respuesta']}:{tol}{expl}}}")
        else:
            ops = p["opciones"]
            if tipo == "opcion":
                cuerpo = [("=" if o.get("correcta") else "~") + gift_esc(o["texto"])
                          + (f"#{gift_esc(o['retro'])}" if o.get("retro") else "") for o in ops]
            else:
                n_ok = sum(bool(o.get("correcta")) for o in ops)
                n_mal = len(ops) - n_ok
                peso_ok = _peso(100 / n_ok)
                peso_mal = _peso(-100 / n_mal) if n_mal else "0"
                cuerpo = [f"~%{peso_ok if o.get('correcta') else peso_mal}%" + gift_esc(o["texto"])
                          + (f"#{gift_esc(o['retro'])}" if o.get("retro") else "") for o in ops]
            lineas.append(cab + " {")
            lineas.extend("\t" + c for c in cuerpo)
            if expl:
                lineas.append("\t" + expl)
            lineas.append("}")
        lineas.append("")
    return "\n".join(lineas)


def _peso(x: float) -> str:
    """Moodle solo acepta ciertos porcentajes; usamos el más cercano válido."""
    validos = [100, 90, 83.33333, 80, 75, 70, 66.66667, 60, 50, 40, 33.33333,
               30, 25, 20, 16.66667, 14.28571, 12.5, 11.11111, 10, 5]
    signo = -1 if x < 0 else 1
    v = min(validos, key=lambda w: abs(w - abs(x)))
    return f"{signo * v:g}"


# --------------------------------------------------------------------------
# Moodle XML
# --------------------------------------------------------------------------
def _texto_xml(padre: ET.Element, etiqueta: str, contenido: str, formato: str = "html") -> None:
    nodo = ET.SubElement(padre, etiqueta, format=formato)
    ET.SubElement(nodo, "text").text = a_html(contenido)


def a_moodle_xml(banco: dict) -> bytes:
    quiz = ET.Element("quiz")
    cat = ET.SubElement(quiz, "question", type="category")
    c = ET.SubElement(cat, "category")
    ET.SubElement(c, "text").text = f"$course$/Reconocimiento de Patrones/{banco['titulo']}"

    for p in banco["preguntas"]:
        tipo = p.get("tipo", "opcion")
        tipo_moodle = {"opcion": "multichoice", "multiple": "multichoice",
                       "vf": "truefalse", "numerica": "numerical"}[tipo]
        q = ET.SubElement(quiz, "question", type=tipo_moodle)
        nombre = ET.SubElement(q, "name")
        ET.SubElement(nombre, "text").text = p["id"]
        _texto_xml(q, "questiontext", p["enunciado"])
        _texto_xml(q, "generalfeedback", p.get("explicacion", ""))
        ET.SubElement(q, "defaultgrade").text = "1"

        if tipo in {"opcion", "multiple"}:
            ET.SubElement(q, "single").text = "true" if tipo == "opcion" else "false"
            ET.SubElement(q, "shuffleanswers").text = "true"
            ET.SubElement(q, "answernumbering").text = "abc"
            ops = p["opciones"]
            n_ok = sum(bool(o.get("correcta")) for o in ops)
            n_mal = len(ops) - n_ok
            for o in ops:
                if tipo == "opcion":
                    frac = "100" if o.get("correcta") else "0"
                else:
                    frac = _peso(100 / n_ok) if o.get("correcta") else (_peso(-100 / n_mal) if n_mal else "0")
                a = ET.SubElement(q, "answer", fraction=frac, format="html")
                ET.SubElement(a, "text").text = a_html(o["texto"])
                fb = ET.SubElement(a, "feedback", format="html")
                ET.SubElement(fb, "text").text = a_html(o.get("retro", ""))
        elif tipo == "vf":
            for valor, etiqueta in ((True, "true"), (False, "false")):
                a = ET.SubElement(q, "answer", fraction="100" if p["respuesta"] is valor else "0",
                                  format="moodle_auto_format")
                ET.SubElement(a, "text").text = etiqueta
        else:
            a = ET.SubElement(q, "answer", fraction="100", format="moodle_auto_format")
            ET.SubElement(a, "text").text = str(p["respuesta"])
            ET.SubElement(a, "tolerance").text = str(p.get("tolerancia", 0))

    ET.indent(quiz)
    return ET.tostring(quiz, encoding="utf-8", xml_declaration=True)


# --------------------------------------------------------------------------
# DOCX para Microsoft Forms (Quick Import)
# --------------------------------------------------------------------------
def a_docx(banco: dict, ruta: Path) -> None:
    try:
        from docx import Document
    except ImportError:  # pragma: no cover
        print("  (omito DOCX: instala python-docx con `uv sync --group dev`)", file=sys.stderr)
        return
    doc = Document()
    doc.add_heading(banco["titulo"], level=1)
    clave: list[tuple[str, str]] = []
    for n, p in enumerate(banco["preguntas"], 1):
        tipo = p.get("tipo", "opcion")
        doc.add_paragraph(f"{n}. {plano(p['enunciado'])}")
        if tipo in {"opcion", "multiple"}:
            letras = "abcdefghij"
            for letra, o in zip(letras, p["opciones"]):
                doc.add_paragraph(f"{letra}. {plano(o['texto'])}")
            ok = ", ".join(l for l, o in zip(letras, p["opciones"]) if o.get("correcta"))
        elif tipo == "vf":
            doc.add_paragraph("a. Verdadero")
            doc.add_paragraph("b. Falso")
            ok = "a" if p["respuesta"] else "b"
        else:
            ok = f"{p['respuesta']} (± {p.get('tolerancia', 0)})"
        clave.append((str(n), ok))
        doc.add_paragraph("")
    doc.add_page_break()
    doc.add_heading("Clave de respuestas (no importar)", level=2)
    tabla = doc.add_table(rows=1, cols=2)
    tabla.rows[0].cells[0].text, tabla.rows[0].cells[1].text = "Pregunta", "Respuesta"
    for n, ok in clave:
        fila = tabla.add_row().cells
        fila[0].text, fila[1].text = n, ok
    doc.save(ruta)


# --------------------------------------------------------------------------
def main(argv: list[str] | None = None) -> int:
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("bancos", nargs="*", help="nombres como u01 (por omisión, todos)")
    ap.add_argument("--validar", action="store_true", help="solo validar, sin exportar")
    args = ap.parse_args(argv)

    nombres = args.bancos or sorted(p.stem for p in BANCO.glob("u*.yml"))
    SALIDA.mkdir(parents=True, exist_ok=True)
    fallos = 0
    for nombre in nombres:
        try:
            banco = cargar(nombre)
        except ErrorBanco as e:
            print(f"✗ {nombre}\n{e}", file=sys.stderr)
            fallos += 1
            continue
        n = len(banco["preguntas"])
        if args.validar:
            print(f"✓ {nombre}: {n} preguntas válidas")
            continue
        (SALIDA / f"{nombre}.gift").write_text(a_gift(banco), encoding="utf-8")
        (SALIDA / f"{nombre}.xml").write_bytes(a_moodle_xml(banco))
        a_docx(banco, SALIDA / f"{nombre}.docx")
        print(f"✓ {nombre}: {n} preguntas → gift, xml, docx")
    return 1 if fallos else 0


if __name__ == "__main__":
    sys.exit(main())
